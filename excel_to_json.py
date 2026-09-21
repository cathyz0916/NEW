#!/usr/bin/env python3
"""
excel_to_json.py  –  Converts EA RISE Impact KPI Excel workbook to data.json.
Supports both Fixed Baseline and Dynamic Baseline sheets, matching the full schema
expected by the dashboard (including bkdown, all_rates, sums, accs).
"""

import json
import os
import sys
from collections import defaultdict
from datetime import date as date_type, datetime
import openpyxl

EXCEL_PATH = "data/latest.xlsx"
OUT_PATH   = "data/data.json"

# Column indices (0-based)
# Col 0: Date (YYYY-MM-DD string or Excel date)
# Col 1: Region (ANZ / G.China / India / Japan / Korea / South East Asia)
# Col 2: Sub-region (e.g. SL-ANZ-AU-NSW)
# Col 3: Customer Name
# Col 4: Customer ID
# Col 5: EA Name
# Col 29: EXIT_ACV (cACV)

MEASURES = [
    {"id": "compliance_check",  "name": "Compliance Check",       "short": "Compliance", "bc": None, "col": 6},
    {"id": "qbr_govnc",         "name": "QBR Governance",          "short": "QBR Gvnc",  "bc": None, "col": 7},
    {"id": "s4_upgrade",        "name": "S4 Upgrade Commitment",   "short": "S4 Upgrd",  "bc": 8,    "col": 9},
    {"id": "ai_adoption_plan",  "name": "AI Adoption Plan",        "short": "AI Adopt",  "bc": None, "col": 10},
    {"id": "ai_agent",          "name": "AI Agent Activated",      "short": "AI Agent",  "bc": 11,   "col": 12},
    {"id": "ai_btp",            "name": "AI BTP Services",         "short": "AI BTP",    "bc": 11,   "col": 14},
    {"id": "calm_tenant",       "name": "CALM Tenant",             "short": "CALM",      "bc": None, "col": 20},
    {"id": "churn_mitigation",  "name": "Churn Mitigation",        "short": "Churn",     "bc": None, "col": 16},
    {"id": "joule",             "name": "Joule Activated",         "short": "Joule",     "bc": 11,   "col": 13},
    {"id": "lead_created",      "name": "Lead Created",            "short": "Lead",      "bc": None, "col": 17},
    {"id": "leanix",            "name": "LeanIX Usage",            "short": "LeanIX",    "bc": 23,   "col": 24},
    {"id": "risk_assessment",   "name": "Risk Assessment",         "short": "Risk Asmt", "bc": None, "col": 15},
    {"id": "rwsm_dibo",         "name": "RwSM DiBO Activated",     "short": "RwSM DiBO", "bc": None, "col": 18},
    {"id": "signavio",          "name": "Signavio Usage",          "short": "Signavio",  "bc": 21,   "col": 22},
    {"id": "tricentis",         "name": "Tricentis",               "short": "Tricentis", "bc": None, "col": 25},
    {"id": "walkme",            "name": "WalkMe Usage",            "short": "WalkMe",    "bc": 26,   "col": 27},
    {"id": "cc_princ_score",    "name": "CC Princ Score",          "short": "CC Score",  "bc": None, "col": 19},
]

CATS = [
    {"id": "ai",        "label": "AI",         "ids": ["ai_adoption_plan","ai_agent","ai_btp","joule"]},
    {"id": "cleancore", "label": "Clean Core",  "ids": ["cc_princ_score","rwsm_dibo"]},
    {"id": "eakpi",     "label": "EA KPI",      "ids": ["churn_mitigation","compliance_check","lead_created","qbr_govnc"]},
    {"id": "rise",      "label": "RISE",        "ids": ["risk_assessment","s4_upgrade"]},
    {"id": "toolchain", "label": "Toolchain",   "ids": ["calm_tenant","leanix","signavio","tricentis","walkme"]},
]

SUB_REGIONS = ["ANZ", "G.China", "India", "Japan", "Korea", "South East Asia"]
BL_I = {8}
BL_C = [8, 11, 21, 23, 26]

REGION_MAP = {
    "anz": "ANZ", "australia": "ANZ", "new zealand": "ANZ", "nz": "ANZ",
    "china": "G.China", "greater china": "G.China", "gchina": "G.China", "g.china": "G.China",
    "sea": "South East Asia", "southeast asia": "South East Asia", "south east asia": "South East Asia",
    "japan": "Japan", "korea": "Korea", "india": "India",
}

def is_true(v):
    if v is None:
        return False
    return str(v).strip().upper() in ("YES", "TRUE", "1", "Y")

def is_false(v):
    if v is None:
        return False
    return str(v).strip().upper() in ("FALSE", "NO", "0", "N")

def fmt_date(v):
    if v is None:
        return None
    if isinstance(v, (date_type, datetime)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return s[:10] if len(s) >= 10 else s

def parse_date(s):
    return datetime.strptime(str(s).strip()[:10], "%Y-%m-%d")

def calc_r(accs, m):
    bc = m["bc"]
    tot = 0
    cnt = 0
    for d in accs.values():
        in_bl = (bc is None or d["b"].get(bc) == 1)
        if in_bl:
            tot += 1
            if d["m"].get(m["id"]) == 1:
                cnt += 1
    if tot == 0:
        return {"rate": None, "count": cnt, "baseline": tot}
    return {"rate": round(cnt / tot * 100, 1), "count": cnt, "baseline": tot}

def agg_r(m, ts, ad):
    bc = m["bc"]
    tot = 0
    cnt = 0
    for reg in SUB_REGIONS:
        accs = ad.get(reg, {}).get(ts, {})
        for d in accs.values():
            in_bl = (bc is None or d["b"].get(bc) == 1)
            if in_bl:
                tot += 1
                if d["m"].get(m["id"]) == 1:
                    cnt += 1
    if tot == 0:
        return {"rate": None, "count": cnt, "baseline": tot}
    return {"rate": round(cnt / tot * 100, 1), "count": cnt, "baseline": tot}

def mk_e(c, w, p):
    wd = round(c["rate"] - w["rate"], 1) if (c["rate"] is not None and w["rate"] is not None) else None
    md = round(c["rate"] - p["rate"], 1) if (c["rate"] is not None and p["rate"] is not None) else None
    trend = "na" if wd is None else ("improved" if wd > 0 else ("declined" if wd < 0 else "flat"))
    return {"current": c, "prev_wk": w, "prev_mo": p, "wd": wd, "md": md, "trend": trend}

def bkst(cd, wd2, pd2, mid, bc):
    in_bl = (bc is None or cd["b"].get(bc) == 1)
    if not in_bl:
        return "NB"
    cm = (cd["m"].get(mid) == 1)
    wm = (wd2["m"].get(mid) == 1) if wd2 else False
    mm = (pd2["m"].get(mid) == 1) if pd2 else False
    if cm:
        if not mm:
            return "NM"
        if not wm:
            return "NW"
        return "UC"
    return "DC" if (wm or mm) else "XT"

def process_sheet(ws):
    """Processes a worksheet (Fixed or Dynamic Baseline) and returns full data dict."""
    ad = {reg: defaultdict(dict) for reg in SUB_REGIONS}
    all_ts_set = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        ts = fmt_date(row[0])
        if not ts:
            continue

        raw_reg = str(row[1]).strip() if row[1] is not None else ""
        reg = REGION_MAP.get(raw_reg.lower(), raw_reg)
        if reg not in SUB_REGIONS:
            continue

        sub = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        name = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
        acc_id = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
        ea = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""

        acv = None
        if len(row) > 29 and row[29] is not None:
            v_str = str(row[29]).strip()
            if v_str and v_str != "#":
                try:
                    acv = float(v_str)
                except ValueError:
                    pass

        if acc_id not in ad[reg][ts]:
            ad[reg][ts][acc_id] = {
                "name": name,
                "ea": ea,
                "sub": sub,
                "acv": acv,
                "m": {m["id"]: 0 for m in MEASURES},
                "b": {bc: 0 for bc in BL_C},
            }
        else:
            if ea and not ad[reg][ts][acc_id]["ea"]:
                ad[reg][ts][acc_id]["ea"] = ea

        for m in MEASURES:
            col_idx = m["col"]
            if col_idx < len(row) and is_true(row[col_idx]):
                ad[reg][ts][acc_id]["m"][m["id"]] = 1

        for bc in BL_C:
            if bc < len(row):
                val = is_false(row[bc]) if bc in BL_I else is_true(row[bc])
                if val:
                    ad[reg][ts][acc_id]["b"][bc] = 1

        all_ts_set.add(ts)

    all_ts = sorted(all_ts_set, key=parse_date)
    if not all_ts:
        return None

    latest = all_ts[-1]
    prev_wk = all_ts[-2] if len(all_ts) > 1 else latest

    lat_dt = parse_date(latest)
    prev_mo = all_ts[0]
    for ts in reversed(all_ts[:-1]):
        d = parse_date(ts)
        if d.month != lat_dt.month or d.year != lat_dt.year:
            prev_mo = ts
            break

    # KPI summary
    new_kpis = {}
    for reg in SUB_REGIONS:
        new_kpis[reg] = {}
        for m in MEASURES:
            new_kpis[reg][m["id"]] = mk_e(
                calc_r(ad[reg].get(latest, {}), m),
                calc_r(ad[reg].get(prev_wk, {}), m),
                calc_r(ad[reg].get(prev_mo, {}), m),
            )

    new_kpis["APAC"] = {}
    for m in MEASURES:
        new_kpis["APAC"][m["id"]] = mk_e(
            agg_r(m, latest, ad),
            agg_r(m, prev_wk, ad),
            agg_r(m, prev_mo, ad),
        )

    # Sums
    new_sums = {}
    for reg in ["APAC"] + SUB_REGIONS:
        s = {"total": len(MEASURES), "improved": 0, "declined": 0, "flat": 0, "na": 0}
        for m in MEASURES:
            t = new_kpis[reg][m["id"]]["trend"]
            if t in s:
                s[t] += 1
        new_sums[reg] = s

    # Accs (for latest, prev_wk, prev_mo)
    new_accs = {}
    for reg in SUB_REGIONS:
        new_accs[reg] = {}
        for ts in (latest, prev_wk, prev_mo):
            new_accs[reg][ts] = {
                acc_id: {"name": d["name"], "ea": d["ea"], "m": d["m"], "b": d["b"]}
                for acc_id, d in ad[reg].get(ts, {}).items()
            }

    new_accs["APAC"] = {}
    for ts in (latest, prev_wk, prev_mo):
        new_accs["APAC"][ts] = {}
        for reg in SUB_REGIONS:
            for acc_id, d in ad[reg].get(ts, {}).items():
                new_accs["APAC"][ts][f"{reg}::{acc_id}"] = {
                    "name": d["name"], "ea": d["ea"], "region": reg, "m": d["m"], "b": d["b"]
                }

    # Breakdown list
    new_bkdown = []
    for reg in SUB_REGIONS:
        latest_accs = ad[reg].get(latest, {})
        pwd = ad[reg].get(prev_wk, {})
        pmd = ad[reg].get(prev_mo, {})
        for acc_id, cd in latest_accs.items():
            new_bkdown.append({
                "account": cd["name"],
                "acc_id": acc_id,
                "ea": cd["ea"] or "",
                "region": reg,
                "sub": cd["sub"] or "",
                "acv": cd["acv"],
                "kpis": {m["id"]: bkst(cd, pwd.get(acc_id), pmd.get(acc_id), m["id"], m["bc"]) for m in MEASURES},
            })
    new_bkdown.sort(key=lambda x: (x["region"], x["account"].lower()))

    # All Rates
    new_all_rates = {}
    for ts in all_ts:
        new_all_rates[ts] = {}
        for reg in SUB_REGIONS:
            new_all_rates[ts][reg] = {
                m["id"]: calc_r(ad[reg].get(ts, {}), m)
                for m in MEASURES
            }
        new_all_rates[ts]["APAC"] = {
            m["id"]: agg_r(m, ts, ad)
            for m in MEASURES
        }

    return {
        "ts": {"latest": latest, "prev_wk": prev_wk, "prev_mo": prev_mo, "all": all_ts},
        "kpis": new_kpis,
        "sums": new_sums,
        "accs": new_accs,
        "bkdown": new_bkdown,
        "all_rates": new_all_rates,
    }

def main():
    if not os.path.exists(EXCEL_PATH):
        print(f"ERROR: File not found: {EXCEL_PATH}")
        sys.exit(1)

    print(f"Loading {EXCEL_PATH} ...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet_names = wb.sheetnames
    print(f"Sheets found: {sheet_names}")

    # Process Fixed Baseline
    fixed_sheet_name = None
    for name in ["Impact KPI_Fixed Baseline", "Fixed Baseline", "Fixed", "KPI", "Sheet1"]:
        if name in sheet_names:
            fixed_sheet_name = name
            break

    if not fixed_sheet_name:
        fixed_sheet_name = sheet_names[0]

    print(f"Parsing Fixed Baseline sheet: {fixed_sheet_name!r}")
    fixed_data = process_sheet(wb[fixed_sheet_name])
    if not fixed_data:
        print(f"ERROR: No valid data found in {fixed_sheet_name}")
        sys.exit(1)

    print(f"  Fixed latest: {fixed_data['ts']['latest']} ({len(fixed_data['bkdown'])} accounts)")

    # Process Dynamic Baseline
    dynamic_sheet_name = None
    for name in ["Impact KPI_Dynamic Baseline", "Dynamic Baseline", "Dynamic"]:
        if name in sheet_names:
            dynamic_sheet_name = name
            break

    if dynamic_sheet_name:
        print(f"Parsing Dynamic Baseline sheet: {dynamic_sheet_name!r}")
        dynamic_data = process_sheet(wb[dynamic_sheet_name])
        print(f"  Dynamic latest: {dynamic_data['ts']['latest']} ({len(dynamic_data['bkdown'])} accounts)")
    else:
        print("  Notice: Dynamic Baseline sheet not found. Falling back to clone of Fixed Baseline.")
        dynamic_data = fixed_data

    DATA = {
        "measures": MEASURES,
        "cats": CATS,
        "regions": ["APAC"] + SUB_REGIONS,
        "fixed": fixed_data,
        "dynamic": dynamic_data,
    }

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(DATA, f, separators=(",", ":"), ensure_ascii=False)

    size_kb = os.path.getsize(OUT_PATH) / 1024
    print(f"\nWritten: {OUT_PATH}  ({size_kb:.1f} KB)")
    print(f"Finished successfully.")

if __name__ == "__main__":
    main()
